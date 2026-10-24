"""Build ``outputs/logistics_and_network_design.xlsx`` (15 sheets, visual-first).

Every number comes from module outputs in ``outputs/logistics/tables``; the
workbook adds live formulas (deltas, national roll-ups, the facility NPV model
and target status) on top. Run after the analytical modules:

    python -m src.logistics.reporting.excel_report
"""
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties

from src.shared.config import load_config, resolve_path
from src.shared.io_utils import load_json
from src.shared.logging_utils import get_logger

from ..common import load_bundle, minutes_to_clock
from ..data_generator import master_register
from .xl_helpers import (
    FMT,
    INK,
    LEAF,
    MARIGOLD,
    SIGNAL,
    SLATE,
    TEAL,
    bar_chart,
    col_range,
    color_scale,
    data_bar,
    font,
    heat_scale,
    inject_sparklines,
    kpi_card,
    line_chart,
    note,
    section,
    setup_sheet,
    write_df,
)

log = get_logger(__name__)
SHEETS = ["Master Data & Assumptions", "Depot & Vehicle Master", "Order & Demand Profile", "Routing Analysis",
          "Route Stops & Compliance", "Fleet Sizing", "Warehouse Slotting", "Warehouse Performance", "Simulation Results",
          "Network Design", "Facility Economics", "ML Performance", "KPI Definitions", "Scenario Comparison",
          "Executive Dashboard"]
REQUIRED = ["routing_summary.json", "routing_city_summary.parquet", "routes_summary.parquet", "route_stops.parquet",
            "unassigned_orders.parquet", "fleet_scenarios.parquet", "fleet_tco.parquet", "warehouse_summary.json",
            "warehouse_performance.parquet", "warehouse_heatmap.parquet", "warehouse_zone_summary.parquet",
            "simulation_summary.parquet", "simulation_cycle_hist.parquet", "network_scenarios.parquet",
            "network_city_ranking.parquet", "network_sensitivity.parquet", "network_sites.parquet",
            "ml_performance.parquet", "ml_zone_factors.parquet", "ml_feature_importance.parquet",
            "ml_demand_validation.parquet", "ml_demand_forecast.parquet", "command_center_view.json", "cc_alerts.parquet"]
GREEN = "008000"


def q(sheet: str) -> str:
    return f"'{sheet}'"


class WorkbookBuilder:
    """Holds loaded results, the workbook and cross-sheet cell references."""

    def __init__(self, cfg: dict[str, Any]) -> None:
        self.cfg = cfg
        self.t = resolve_path(cfg, "tables_dir")
        missing = [f for f in REQUIRED if not (self.t / f).exists()]
        if missing:
            raise FileNotFoundError(f"Run the analytical modules first; missing outputs: {missing}")
        self.b = load_bundle(cfg)
        self.names = {k: v["name"] for k, v in cfg["cities"].items()}
        self.cities = list(cfg["cities"])
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws = {s: self.wb.create_sheet(s) for s in SHEETS}
        self.ref: dict[str, str] = {}
        self.spark: dict[str, list[tuple[str, str]]] = {}

    def p(self, name: str) -> pd.DataFrame:
        return pd.read_parquet(self.t / name)

    def j(self, name: str) -> Any:
        return load_json(self.t / name)

    # ------------------------------------------------------------------ 1
    def master(self) -> None:
        s = SHEETS[0]
        ws = self.ws[s]
        setup_sheet(ws, "Logistics and Network Design",
                    "Portfolio Output 6 · five metro clusters across seven states/UTs · 90-day operating history "
                    f"from {self.cfg['calendar']['start_date']}", 14, INK)
        o, a = self.b.orders, self.b.addresses
        cards = [("Metro clusters", len(self.cities), "int"), ("States / UTs", a["state"].nunique(), "int"),
                 ("Delivery addresses", len(a), "int"), ("Vehicles", len(self.b.vehicles), "int"),
                 ("Orders (90 days)", len(o), "int"), ("Candidate MFC sites", len(self.b.candidate_sites), "int")]
        for k, (lab, val, f) in enumerate(cards):
            kpi_card(ws, 4, 2 + 2 * k, lab, int(val), FMT[f])
        section(ws, 7, 2, "City cluster profile", 12)
        days = o["order_date"].nunique()
        rows = []
        for cid, c in self.cfg["cities"].items():
            oc = o[o["city_id"] == cid]
            daily = oc.groupby("order_date").size()
            rows.append({"City ID": cid, "City cluster": c["name"],
                         "States / UTs": ", ".join(sorted(a.loc[a["city_id"] == cid, "state"].unique())),
                         "Volume weight": c["volume_weight"], "Congestion index": c["congestion"], "Wage index": c["wage_index"],
                         "Orders / day": round(len(oc) / days, 1), "Peak-day orders": int(daily.max()),
                         "Peak date": daily.idxmax().date(), "E-com share": round((oc["segment"] == "ECOM_B2C").mean(), 4),
                         "Avg kg / order": round(oc["weight_kg"].mean(), 1)})
        prof = pd.DataFrame(rows)
        r0, r1, c0, c1 = write_df(ws, prof, 8, 2, {"Volume weight": "dec2", "Congestion index": "dec2", "Wage index": "dec2",
                                                    "Orders / day": "dec1", "Peak-day orders": "int", "E-com share": "pct",
                                                    "Avg kg / order": "dec1"}, "CityProfile",
                                  {"City cluster": 16, "States / UTs": 30})
        data_bar(ws, col_range(r0 + 1, r1, c0 + 6))
        color_scale(ws, col_range(r0 + 1, r1, c0 + 4), reverse=True)
        bar_chart(ws, f"B{r1 + 2}", ws, (c0 + 1, r0 + 1, r1), [c0 + 6, c0 + 7], r0, "Orders per day vs peak day by city",
                  "orders", w=17, h=7)
        # Editable levers (blue inputs) used by Facility Economics formulas
        lr = r1 + 2
        section(ws, lr, 11, "Scenario levers (editable)", 4, "Blue = input; Facility Economics recalculates from these cells.")
        net = self.cfg["network"]
        levers = [("WACC", net["wacc"], "pct"), ("Demand growth / year", net["demand_growth"], "pct"),
                  ("Compliance target", self.cfg["fleet_sizing"]["compliance_target"], "pct"),
                  ("P95 cycle-time target (min)", self.cfg["simulation"]["cycle_time_target_min"], "int")]
        for k, (lab, val, f) in enumerate(levers):
            r = lr + 2 + k
            ws.cell(r, 11, lab).font = font(10)
            c = ws.cell(r, 13, val)
            c.font = font(10, True, "0000FF")
            c.number_format = FMT[f]
            c.fill = PatternFill("solid", fgColor="FFF6DD")
            self.ref[lab] = f"{q(s)}!$M${r}"
        reg_row = r1 + 18
        section(ws, reg_row, 2, "Master data register", 12,
                "Every parameter used by the models, with unit, rationale and dependent modules (config/logistics.yaml).")
        reg = master_register(self.cfg)
        reg.columns = ["Parameter", "Value", "Unit", "Rationale", "Dependent modules"]
        reg["Value"] = reg["Value"].astype(str).str.slice(0, 120)
        write_df(ws, reg, reg_row + 2, 2, None, "MasterRegister", {})
        ws.column_dimensions["B"].width = 34
        note(ws, reg_row + 3 + len(reg), 2, "Colour legend: blue font = editable input; black = formula or model output; "
             "green = link to another sheet. All figures are produced by the Python modules from the operational dataset.", 12)
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ 2
    def depots(self) -> None:
        s = SHEETS[1]
        ws = self.ws[s]
        setup_sheet(ws, "Depot & Vehicle Master", "15 depots (DC, MFC, hub per city) · 125 vehicles · vehicle economics", 16)
        d = self.b.depots.copy()
        vcount = self.b.vehicles.groupby("depot_id").size()
        d["Vehicles"] = d["depot_id"].map(vcount).fillna(0).astype(int)
        d["Opens"] = d["open_min"].map(minutes_to_clock)
        d["Closes"] = d["close_min"].map(minutes_to_clock)
        d = d[["depot_id", "city", "state", "role", "name", "zone", "lat", "lon", "Vehicles", "Opens", "Closes"]]
        d.columns = ["Depot ID", "City", "State", "Role", "Depot", "Zone", "Lat", "Lon", "Vehicles", "Opens", "Closes"]
        section(ws, 4, 2, "Depot master", 11)
        r0, r1, c0, _ = write_df(ws, d, 5, 2, {"Lat": "0.0000", "Lon": "0.0000"}, "Depots", {"Depot": 22, "State": 24})
        data_bar(ws, col_range(r0 + 1, r1, c0 + 8), TEAL)
        vt = self.b.vehicle_types[["vehicle_type", "label", "capacity_kg", "speed_kmph", "cost_per_km", "fixed_cost_per_day",
                                   "driver_wage_per_hour", "max_route_minutes", "max_order_kg", "acquisition_cost",
                                   "lease_per_month", "restricted_zones"]].copy()
        vt.columns = ["Type", "Label", "Capacity kg", "Speed km/h", "₹ / km", "Fixed ₹ / day", "Driver ₹ / h (MUM)",
                      "Duty min", "Max parcel kg", "Acquisition ₹", "Lease ₹ / month", "Restricted zones"]
        vr = r1 + 3
        section(ws, vr, 2, "Vehicle types", 11, "Driver wages are scaled by each city's wage index in the vehicle master.")
        r0, r1, c0, _ = write_df(ws, vt, vr + 2, 2, {"₹ / km": "inr2", "Fixed ₹ / day": "inr", "Driver ₹ / h (MUM)": "inr",
                                                     "Acquisition ₹": "inr", "Lease ₹ / month": "inr", "Capacity kg": "int"},
                                 "VehicleTypes", {"Restricted zones": 60})
        pv = self.b.vehicles.pivot_table(index="city_id", columns="vehicle_type", values="vehicle_id", aggfunc="count").fillna(0)
        pv = pv.reindex(self.cities).reset_index()
        pv.insert(1, "City", pv["city_id"].map(self.names))
        pv = pv.drop(columns="city_id")
        fr = r1 + 3
        section(ws, fr, 2, "Fleet by city and type (current)", 6)
        r0, r1, c0, c1 = write_df(ws, pv, fr + 2, 2, {c: "int" for c in pv.columns[1:]}, "FleetByCity")
        bar_chart(ws, f"I{fr}", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1 + 1)), r0, "Current fleet by city and vehicle type",
                  "vehicles", stacked=True, colors=[INK, TEAL, MARIGOLD], w=15, h=7)
        econ_r = r1 + 12
        bar_chart(ws, f"B{econ_r}", ws, (2, vr + 3, vr + 5), [6, 7], vr + 2, "Running vs fixed cost by vehicle type",
                  "₹", colors=[TEAL, SLATE], w=15, h=7)
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ 3
    def demand(self) -> None:
        s = SHEETS[2]
        ws = self.ws[s]
        setup_sheet(ws, "Order & Demand Profile", "Daily volume by city with national and regional events, segment mix and weekly heat", 18)
        o, cal = self.b.orders, self.b.calendar
        daily = o.pivot_table(index="order_date", columns="city_id", values="order_id", aggfunc="count").fillna(0)
        daily = daily[self.cities]
        daily.columns = [self.names[c] for c in self.cities]
        daily["All cities"] = daily.sum(axis=1)
        daily = daily.reset_index().rename(columns={"order_date": "Date"})
        normal = cal.groupby("date")["is_event"].any()
        peak_days = cal[cal["event"].str.contains("Diwali peak")]["date"].unique()
        tot = daily.set_index("Date")["All cities"]
        uplift = tot.loc[peak_days].mean() / tot[~normal.reindex(tot.index).values].mean()
        cards = [("Orders, 90 days", len(o), "int"), ("Average orders / day", tot.mean(), "dec1"),
                 (f"Peak day ({tot.idxmax().date()})", tot.max(), "int"), ("Diwali-peak uplift vs normal days", uplift, "x")]
        for k, (lab, val, f) in enumerate(cards):
            kpi_card(ws, 4, 2 + 3 * k, lab, float(val), FMT[f], span=3, accent=MARIGOLD if k == 3 else TEAL)
        section(ws, 7, 2, "Daily orders by city", 8)
        r0, r1, c0, c1 = write_df(ws, daily, 8, 2, {c: "int" for c in daily.columns[1:]}, "DailyOrders", {"Date": 12})
        line_chart(ws, "K8", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1)), r0,
                   "Daily orders by city: festive sale, Navratri, Diwali peak and holidays", "orders/day",
                   colors=["1F2A44", "0E7C86", "E3A21A", "8E5A9B", "C2413B"], w=24, h=9)
        # Sparkline trend table
        sr = 28
        section(ws, sr, 11, "Trend by city", 7)
        hdr = ["City", "90-day trend", "Min", "Average", "Max", "Peak date"]
        for k, h in enumerate(hdr):
            c = ws.cell(sr + 1, 11 + k, h)
            c.font = font(10, True, "FFFFFF")
            c.fill = PatternFill("solid", fgColor=INK)
        self.spark[s] = []
        for k, name in enumerate(list(daily.columns[1:])):
            r = sr + 2 + k
            col = c0 + 1 + k
            rng = f"{L(col)}{r0 + 1}:{L(col)}{r1}"
            ws.cell(r, 11, name).font = font(10, True)
            for j, fn in enumerate(("MIN", "AVERAGE", "MAX")):
                c = ws.cell(r, 13 + j, f"={fn}({rng})")
                c.number_format = FMT["dec1"] if fn == "AVERAGE" else FMT["int"]
                c.font = font(10)
            c = ws.cell(r, 16, f"=INDEX($B${r0 + 1}:$B${r1},MATCH(MAX({rng}),{rng},0))")
            c.number_format = "dd-mmm"
            self.spark[s].append((f"{q(s)}!{rng}", f"L{r}"))
            ws.row_dimensions[r].height = 22
        ws.column_dimensions["L"].width = 26
        # Segment mix
        mr = sr + 10
        section(ws, mr, 11, "Segment mix by city (share of orders)", 7)
        mix = (o.pivot_table(index="city_id", columns="segment", values="order_id", aggfunc="count")
               .div(o.groupby("city_id").size(), axis=0).reindex(self.cities).reset_index())
        mix.insert(0, "City", mix["city_id"].map(self.names))
        mix = mix.drop(columns="city_id")
        r0m, r1m, c0m, c1m = write_df(ws, mix, mr + 1, 11, {c: "pct" for c in mix.columns[1:]}, "SegmentMix")
        bar_chart(ws, f"K{r1m + 2}", ws, (c0m, r0m + 1, r1m), list(range(c0m + 1, c1m + 1)), r0m,
                  "Segment mix by city", "share", stacked=True, w=16, h=7, number_format="0%")
        # Weekly heat
        wr = r1 + 3
        section(ws, wr, 2, "Weekly orders heatmap (city x ISO week)", 9)
        o2 = o.assign(week=o["order_date"].dt.isocalendar().week.astype(int))
        heat = o2.pivot_table(index="city_id", columns="week", values="order_id", aggfunc="count").reindex(self.cities)
        heat.columns = [f"W{w}" for w in heat.columns]
        heat = heat.reset_index()
        heat.insert(0, "City", heat["city_id"].map(self.names))
        heat = heat.drop(columns="city_id")
        r0h, r1h, c0h, c1h = write_df(ws, heat, wr + 1, 2, {c: "int" for c in heat.columns[1:]}, None)
        heat_scale(ws, f"{L(c0h + 1)}{r0h + 1}:{L(c1h)}{r1h}")
        er = r1h + 3
        section(ws, er, 2, "Event calendar (national and regional)", 9)
        evs = [{"Scope": "National", "Event": e["name"], "Start": e["start"], "End": e["end"], "B2C x": e["b2c"],
                "B2B x": e["b2b"], "Traffic x": e["traffic"]} for e in self.cfg["calendar"]["events"]]
        for c in self.cfg["cities"].values():
            evs += [{"Scope": c["name"], "Event": e["name"], "Start": e["start"], "End": e["end"], "B2C x": e["b2c"],
                     "B2B x": e["b2b"], "Traffic x": e["traffic"]} for e in c.get("events", [])]
        ev = pd.DataFrame(evs)
        r0e, r1e, c0e, _ = write_df(ws, ev, er + 1, 2, {"B2C x": "dec2", "B2B x": "dec2", "Traffic x": "dec2"}, "Events",
                                    {"Event": 26})
        color_scale(ws, col_range(r0e + 1, r1e, c0e + 4))
        ws.freeze_panes = "C9"

    # ------------------------------------------------------------------ 4
    def routing(self) -> None:
        s = SHEETS[3]
        ws = self.ws[s]
        rs = self.j("routing_summary.json")
        setup_sheet(ws, "Routing Analysis", f"CVRPTW (OR-Tools, guided local search) vs window-aware nearest-neighbour baseline · "
                    f"representative day {rs['day']} · {rs['orders']:,} orders", 17)
        cm = self.p("routing_city_summary.parquet")
        b = cm[cm["plan"] == "baseline"].set_index("city_id").reindex(self.cities)
        o = cm[cm["plan"] == "optimized"].set_index("city_id").reindex(self.cities)
        tbl = pd.DataFrame({"City": [self.names[c] for c in self.cities], "Orders": b["orders"].values,
                            "Baseline last-mile ₹": b["last_mile_cost"].values, "Optimized last-mile ₹": o["last_mile_cost"].values,
                            "Last-mile reduction": None,
                            "Baseline fleet ₹": b["fleet_cost"].values, "Optimized fleet ₹": o["fleet_cost"].values,
                            "Fleet cost reduction": None,
                            "Baseline compliance": b["tw_compliance"].values, "Optimized compliance": o["tw_compliance"].values,
                            "Baseline vehicles": b["vehicles_used"].values, "Optimized vehicles": o["vehicles_used"].values,
                            "Baseline unassigned": b["unassigned"].values, "Optimized unassigned": o["unassigned"].values,
                            "Baseline road km": b["road_km"].values, "Optimized road km": o["road_km"].values})
        section(ws, 7, 2, "Baseline vs optimized by city (formulas compute deltas and national totals)", 16)
        r0, r1, c0, c1 = write_df(ws, tbl, 8, 2, {"Orders": "int", "Baseline last-mile ₹": "inr", "Optimized last-mile ₹": "inr",
                                                   "Last-mile reduction": "pct", "Baseline fleet ₹": "inr", "Optimized fleet ₹": "inr",
                                                   "Fleet cost reduction": "pct", "Baseline compliance": "pct",
                                                   "Optimized compliance": "pct", "Baseline road km": "dec1",
                                                   "Optimized road km": "dec1"}, "RoutingByCity", {"City": 15})
        for r in range(r0 + 1, r1 + 1):
            ws.cell(r, c0 + 4, f"=IFERROR(1-{L(c0 + 3)}{r}/{L(c0 + 2)}{r},0)").number_format = FMT["pct"]
            ws.cell(r, c0 + 7, f"=IFERROR(1-{L(c0 + 6)}{r}/{L(c0 + 5)}{r},0)").number_format = FMT["pct"]
        tr = r1 + 1
        ws.cell(tr, c0, "All cities").font = font(10, True)
        for k in range(1, c1 - c0 + 1):
            col = c0 + k
            letter = L(col)
            if k in (4, 7):
                f = f"=IFERROR(1-{L(col - 1)}{tr}/{L(col - 2)}{tr},0)"
            elif k in (8, 9):
                f = f"=SUMPRODUCT({letter}{r0 + 1}:{letter}{r1},$C${r0 + 1}:$C${r1})/SUM($C${r0 + 1}:$C${r1})"
            else:
                f = f"=SUM({letter}{r0 + 1}:{letter}{r1})"
            c = ws.cell(tr, col, f)
            c.font = font(10, True)
            c.number_format = ws.cell(r1, col).number_format
            c.fill = PatternFill("solid", fgColor="E3EEF0")
        data_bar(ws, col_range(r0 + 1, r1, c0 + 4), TEAL)
        data_bar(ws, col_range(r0 + 1, r1, c0 + 7), LEAF)
        color_scale(ws, col_range(r0 + 1, r1, c0 + 9))
        self.ref["routing_lm_red"] = f"{q(s)}!{L(c0 + 4)}{tr}"
        self.ref["routing_fleet_red"] = f"{q(s)}!{L(c0 + 7)}{tr}"
        self.ref["routing_comp"] = f"{q(s)}!{L(c0 + 9)}{tr}"
        self.ref["routing_comp_base"] = f"{q(s)}!{L(c0 + 8)}{tr}"
        self.ref["veh_base"] = f"{q(s)}!{L(c0 + 10)}{tr}"
        self.ref["veh_opt"] = f"{q(s)}!{L(c0 + 11)}{tr}"
        self.ref["routing_tbl"] = (s, r0, r1, c0)
        kpi_card(ws, 4, 2, "Last-mile cost reduction (incl. outsourcing)", f"={L(c0 + 4)}{tr}", FMT["pct"], 3)
        kpi_card(ws, 4, 5, "Fleet operating cost reduction", f"={L(c0 + 7)}{tr}", FMT["pct"], 3)
        kpi_card(ws, 4, 8, "Optimized in-window delivery", f"={L(c0 + 9)}{tr}", FMT["pct"], 3, LEAF)
        kpi_card(ws, 4, 11, "Vehicles used (optimized)", f"={L(c0 + 11)}{tr}", FMT["int"], 3, sub="vs baseline in table")
        kpi_card(ws, 4, 14, "Road km saved", f"={L(c0 + 14)}{tr}-{L(c0 + 15)}{tr}", FMT["km"], 3)
        cr = tr + 2
        bar_chart(ws, f"B{cr}", ws, (c0, r0 + 1, r1), [c0 + 2, c0 + 3], r0, "Daily last-mile cost by city", "₹ per day",
                  colors=[SLATE, TEAL], w=16, h=8, number_format='"₹"#,##0')
        bar_chart(ws, f"J{cr}", ws, (c0, r0 + 1, r1), [c0 + 8, c0 + 9], r0, "Time-window compliance by city", "share",
                  colors=[SLATE, LEAF], w=16, h=8, number_format="0%")
        bar_chart(ws, f"B{cr + 17}", ws, (c0, r0 + 1, r1), [c0 + 10, c0 + 11], r0, "Vehicles used by city", "vehicles",
                  colors=[SLATE, TEAL], w=16, h=8)
        bar_chart(ws, f"J{cr + 17}", ws, (c0, r0 + 1, r1), [c0 + 14, c0 + 15], r0, "Estimated road km by city", "km",
                  colors=[SLATE, TEAL], w=16, h=8)
        note(ws, cr + 34, 2, "Baseline: window-aware nearest-neighbour dispatch (each vehicle repeatedly takes the closest order it "
             "can reach before its window closes; late second pass). Both plans are replayed by the same evaluator with "
             "departure-time traffic bands and ML zone calibration. Unserved orders are costed at spot-outsourcing rates, so "
             "'last-mile cost' compares like-for-like service. Road km = straight-line km x 1.35 planning circuity "
             "(estimate, not routed street distance).", 16)
        ws.freeze_panes = "C9"

    # ------------------------------------------------------------------ 5
    def stops(self) -> None:
        s = SHEETS[4]
        ws = self.ws[s]
        setup_sheet(ws, "Route Stops & Compliance", "Zone-level service, vehicle utilisation and stop-level audit trail (optimized plan)", 16)
        st = self.p("route_stops.parquet")
        z = st.groupby(["plan", "city_id", "zone"]).agg(stops=("order_id", "size"), on_time=("on_time", "mean"),
                                                         late=("late_min", "mean"), wait=("wait_min", "mean")).reset_index()
        zt = z[z["plan"] == "optimized"].merge(z[z["plan"] == "baseline"][["city_id", "zone", "on_time"]],
                                               on=["city_id", "zone"], suffixes=("", "_b"))
        zt = zt.assign(City=zt["city_id"].map(self.names))[["City", "zone", "stops", "on_time_b", "on_time", "late", "wait"]]
        zt.columns = ["City", "Zone", "Stops", "Baseline on-time", "Optimized on-time", "Avg late min", "Avg wait min"]
        section(ws, 4, 2, "On-time performance by zone", 7)
        r0, r1, c0, _ = write_df(ws, zt, 5, 2, {"Stops": "int", "Baseline on-time": "pct", "Optimized on-time": "pct",
                                                 "Avg late min": "dec1", "Avg wait min": "dec1"}, "ZoneCompliance", {"Zone": 18})
        color_scale(ws, col_range(r0 + 1, r1, c0 + 3))
        color_scale(ws, col_range(r0 + 1, r1, c0 + 4))
        data_bar(ws, col_range(r0 + 1, r1, c0 + 6), MARIGOLD)
        bar_chart(ws, "J4", ws, (c0 + 1, r0 + 1, r1), [c0 + 3, c0 + 4], r0, "On-time rate by zone: baseline vs optimized", "share",
                  colors=[SLATE, TEAL], horizontal=True, w=17, h=16, number_format="0%")
        rt = self.p("routes_summary.parquet")
        rt = rt[rt["plan"] == "optimized"].copy()
        rt["City"] = rt["city_id"].map(self.names)
        rt = rt[["City", "depot_id", "vehicle_id", "vehicle_type", "stops", "load_util_pct", "time_util_pct", "straight_km",
                 "road_km", "duration_min", "wait_min", "late_stops", "duration_violation", "cost_total"]]
        rt.columns = ["City", "Depot", "Vehicle", "Type", "Stops", "Load util %", "Duty util %", "Straight km", "Road km (est.)",
                      "Duration min", "Wait min", "Late stops", "Duty overrun", "Route cost ₹"]
        vr = r1 + 3
        section(ws, vr, 2, "Vehicle routes (optimized)", 14, "Duty overrun = route exceeds driver duty limit when replayed under "
                "actual departure-time traffic; reported as overtime, not hidden.")
        r0, r1, c0, _ = write_df(ws, rt, vr + 2, 2, {"Load util %": "dec1", "Duty util %": "dec1", "Straight km": "dec1",
                                                      "Road km (est.)": "dec1", "Duration min": "dec1", "Wait min": "dec1",
                                                      "Route cost ₹": "inr"}, "VehicleRoutes")
        data_bar(ws, col_range(r0 + 1, r1, c0 + 5), TEAL)
        data_bar(ws, col_range(r0 + 1, r1, c0 + 6), MARIGOLD)
        ws.conditional_formatting.add(col_range(r0 + 1, r1, c0 + 12),
                                      CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="F8D9D6")))
        un = self.p("unassigned_orders.parquet")
        un = un[un["plan"] == "optimized"][["order_id", "city_id", "depot_id", "segment", "zone", "weight_kg", "tw_start_min",
                                            "tw_end_min", "outsourcing_cost", "reason"]]
        ur = r1 + 3
        section(ws, ur, 2, f"Unassigned orders in the optimized plan ({len(un)})", 10,
                "Explicitly reported and costed at the spot-outsourcing rate.")
        if len(un):
            write_df(ws, un, ur + 2, 2, {"outsourcing_cost": "inr"}, "Unassigned", {"reason": 34})
            ur += len(un)
        sr = ur + 5
        stp = st[st["plan"] == "optimized"][["city_id", "vehicle_id", "seq", "order_id", "segment", "zone", "weight_kg",
                                             "tw_start_min", "tw_end_min", "arrival_min", "late_min", "on_time",
                                             "leg_straight_km", "leg_road_km", "leg_travel_min"]].sort_values(["city_id", "vehicle_id", "seq"])
        section(ws, sr, 2, f"Stop-level audit trail ({len(stp):,} stops, filterable)", 14)
        r0, r1, c0, _ = write_df(ws, stp, sr + 1, 2, {"arrival_min": "dec1", "late_min": "dec1", "leg_straight_km": "dec2",
                                                       "leg_road_km": "dec2", "leg_travel_min": "dec1"}, "StopTrail")
        ws.conditional_formatting.add(col_range(r0 + 1, r1, c0 + 11),
                                      CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="F8D9D6")))
        ws.freeze_panes = "C6"

    # ------------------------------------------------------------------ 6
    def fleet(self) -> None:
        s = SHEETS[5]
        ws = self.ws[s]
        setup_sheet(ws, "Fleet Sizing", "Minimum fleet and vehicle mix per demand level (fleet-size-and-mix CVRPTW) · own vs lease TCO", 16)
        fs = self.p("fleet_scenarios.parquet")
        nat = fs[fs["city_id"] == "ALL"].set_index("demand_multiplier")
        one = nat.loc[1.0]
        cur_ok = nat[nat["current_fleet_meets_target"]].index.max()
        kpi_card(ws, 4, 2, "Recommended fleet at 1.0x (all cities)", int(one["recommended_fleet"]), FMT["int"], 3,
                 sub=f"current {int(one['current_fleet_size'])}")
        kpi_card(ws, 4, 5, "Highest level the current fleet serves at target", float(cur_ok) if pd.notna(cur_ok) else 0.0,
                 FMT["x"], 3, MARIGOLD)
        top = float(max(nat.index))   # highest demand level modelled (3.0x in the full run)
        kpi_card(ws, 4, 8, f"Fleet needed at {top:g}x (highest level modelled)", int(nat.loc[top, "recommended_fleet"]), FMT["int"], 3, SIGNAL)
        kpi_card(ws, 4, 11, "Infeasible city-levels", int((~fs["feasible"]).sum()), FMT["int"], 3, SIGNAL)
        levels = sorted(fs["demand_multiplier"].unique())
        rec = fs[fs["city_id"] != "ALL"].pivot_table(index="demand_multiplier", columns="city", values="recommended_fleet")
        rec = rec[[self.names[c] for c in self.cities]].reset_index().rename(columns={"demand_multiplier": "Demand level"})
        section(ws, 7, 2, "Recommended fleet by demand level", 7)
        r0, r1, c0, c1 = write_df(ws, rec, 8, 2, {"Demand level": "x"} | {c: "int" for c in rec.columns[1:]}, "RecFleet")
        line_chart(ws, "J7", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1 + 1)), r0, "Fleet size needed vs demand level",
                   "vehicles", "demand multiplier", markers=True, w=17, h=8)
        comp = fs[fs["city_id"] != "ALL"].pivot_table(index="demand_multiplier", columns="city", values="current_fleet_compliance")
        comp = comp[[self.names[c] for c in self.cities]]
        comp["Target"] = self.cfg["fleet_sizing"]["compliance_target"]
        comp = comp.reset_index().rename(columns={"demand_multiplier": "Demand level"})
        cr = r1 + 3
        section(ws, cr, 2, "Current 25-vehicle fleet: compliance by demand level", 8)
        r0, r1, c0, c1 = write_df(ws, comp, cr + 1, 2, {"Demand level": "x"} | {c: "pct" for c in comp.columns[1:]}, "CurFleetComp")
        color_scale(ws, f"{L(c0 + 1)}{r0 + 1}:{L(c1 - 1)}{r1}")
        line_chart(ws, f"J{cr + 10}", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1 + 1)), r0,
                   "Current fleet compliance falls as demand rises", "in-window share", "demand multiplier",
                   colors=["1F2A44", "0E7C86", "E3A21A", "8E5A9B", "C2413B", "6B7A90"], dashed={len(self.cities)},
                   number_format="0%", markers=True, w=17, h=8)
        sc = fs[["city", "demand_multiplier", "orders", "history_days_at_level", "current_fleet_compliance",
                 "current_fleet_unassigned", "recommended_fleet", "rec_vans", "rec_tempos", "rec_two_wheelers",
                 "rec_compliance", "rec_cost_per_order", "rec_time_util_pct", "annual_tco_own", "annual_tco_lease", "feasible",
                 "status"]].copy()
        sc.columns = ["City", "Level", "Orders", "History days", "Current-fleet compliance", "Current unassigned", "Rec. fleet",
                      "Vans", "Tempos", "2-wheelers", "Rec. compliance", "₹ / order", "Duty util %", "Annual TCO own ₹",
                      "Annual TCO lease ₹", "Feasible", "Status"]
        tr = r1 + 3
        section(ws, tr, 2, "All scenarios (per city and all cities)", 16)
        r0, r1, c0, c1 = write_df(ws, sc, tr + 1, 2, {"Level": "x", "Orders": "int", "Current-fleet compliance": "pct",
                                                       "Rec. compliance": "pct", "₹ / order": "inr2", "Duty util %": "dec1",
                                                       "Annual TCO own ₹": "inr", "Annual TCO lease ₹": "inr"}, "FleetScenarios",
                                  {"Status": 40})
        data_bar(ws, col_range(r0 + 1, r1, c0 + 6), TEAL)
        ws.conditional_formatting.add(col_range(r0 + 1, r1, c0 + 15), CellIsRule(operator="equal", formula=['"Yes"'],
                                      fill=PatternFill("solid", fgColor="D7EEE0")))
        ws.conditional_formatting.add(col_range(r0 + 1, r1, c0 + 15), CellIsRule(operator="equal", formula=['"No"'],
                                      fill=PatternFill("solid", fgColor="F8D9D6")))
        tco = self.p("fleet_tco.parquet")
        t1 = tco[tco["demand_multiplier"] == 1.0].groupby("city_id")[["capital_eac", "insurance", "maintenance", "fuel", "labour",
                                                                       "tco_own", "tco_lease"]].sum().reindex(self.cities)
        t1.insert(0, "City", [self.names[c] for c in t1.index])
        t1.columns = ["City", "Capital (EAC) ₹", "Insurance ₹", "Maintenance ₹", "Fuel ₹", "Labour ₹", "TCO own ₹", "TCO lease ₹"]
        kr = r1 + 3
        section(ws, kr, 2, "Annual fleet TCO at 1.0x demand (recommended fleet)", 9,
                "Own = equivalent annual capital cost at WACC + insurance + maintenance + fuel + labour; lease = lease + admin + fuel + labour.")
        r0, r1, c0, c1 = write_df(ws, t1, kr + 2, 2, {c: "inr" for c in t1.columns[1:]}, "FleetTCO")
        bar_chart(ws, f"B{r1 + 2}", ws, (c0, r0 + 1, r1), [c0 + 1, c0 + 2, c0 + 3, c0 + 4, c0 + 5], r0,
                  "Annual ownership TCO composition by city", "₹ / year", stacked=True,
                  colors=[INK, SLATE, "8E5A9B", MARIGOLD, TEAL], w=16, h=8, number_format='"₹"#,##0,,"M"')
        bar_chart(ws, f"J{r1 + 2}", ws, (c0, r0 + 1, r1), [c0 + 6, c0 + 7], r0, "Own vs lease (annual TCO)", "₹ / year",
                  colors=[TEAL, MARIGOLD], w=16, h=8, number_format='"₹"#,##0,,"M"')
        _ = levels
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ 7
    def slotting(self) -> None:
        s = SHEETS[6]
        ws = self.ws[s]
        whs = self.j("warehouse_summary.json")
        setup_sheet(ws, "Warehouse Slotting", "10,000-slot DC template per city · velocity zoning + co-pick affinity + ergonomic weight rule", 54)
        for c in range(3, 55):
            ws.column_dimensions[L(c)].width = 3.2
        ws.column_dimensions["B"].width = 16
        nat = whs["national"]
        heavy = sum(v["baseline"]["heavy_slot_violations"] for v in whs["by_city"].values())
        kpi_card(ws, 4, 2, "Picker travel reduction (held-out)", nat["travel_reduction_pct"] / 100, FMT["pct"], 10)
        kpi_card(ws, 4, 12, "Pick throughput gain", nat["throughput_gain_pct"] / 100, FMT["pct"], 10, LEAF)
        kpi_card(ws, 4, 22, "Perfect-information bound", nat["perfect_info_travel_reduction_pct"] / 100, FMT["pct"], 10, MARIGOLD)
        kpi_card(ws, 4, 32, "Heavy-item slot violations removed", heavy, FMT["int"], 10, SIGNAL, sub="now 0")
        hm = self.p("warehouse_heatmap.parquet")
        hm = hm[hm["city_id"] == "MUM"]
        row = 7
        for col, label in (("baseline_picks", "Current layout (category order)"), ("optimized_picks", "Optimized layout")):
            section(ws, row, 2, f"Mumbai DC pick intensity - {label} (rows: bay, columns: aisle; dispatch at top-left)", 50)
            grid = hm.pivot_table(index="bay", columns="aisle", values=col).fillna(0)
            ws.cell(row + 1, 2, "Bay \\ Aisle").font = font(8, True)
            for j, a in enumerate(grid.columns):
                c = ws.cell(row + 1, 3 + j, int(a))
                c.font = font(7, color=SLATE)
                c.alignment = Alignment(horizontal="center")
            for i, (bay, vals) in enumerate(grid.iterrows()):
                ws.cell(row + 2 + i, 2, int(bay)).font = font(7, color=SLATE)
                for j, v in enumerate(vals):
                    c = ws.cell(row + 2 + i, 3 + j, round(float(v), 1))
                    c.number_format = ";;;"
                ws.row_dimensions[row + 2 + i].height = 11
            heat_scale(ws, f"C{row + 2}:{L(2 + grid.shape[1])}{row + 1 + grid.shape[0]}")
            row += grid.shape[0] + 4
        ws.freeze_panes = "C4"

    # ------------------------------------------------------------------ 8
    def wh_perf(self) -> None:
        s = SHEETS[7]
        ws = self.ws[s]
        setup_sheet(ws, "Warehouse Performance", "Held-out pick orders · return routing, batches of 3 · labour cost at city wage level", 16)
        perf = self.p("warehouse_performance.parquet")
        rows = []
        for cid in self.cities:
            g = perf[perf["city_id"] == cid].set_index("layout")
            rows.append({"City": self.names[cid], "Pick lines": int(g.at["baseline", "pick_lines"]),
                         "Travel m/line (current)": g.at["baseline", "travel_m_per_line"],
                         "Travel m/line (optimized)": g.at["optimized", "travel_m_per_line"],
                         "Travel m/line (bound)": g.at["perfect_information", "travel_m_per_line"],
                         "Travel reduction": None, "Picks/h (current)": g.at["baseline", "picks_per_hour"],
                         "Picks/h (optimized)": g.at["optimized", "picks_per_hour"],
                         "Picks/h (bound)": g.at["perfect_information", "picks_per_hour"], "Throughput gain": None,
                         "₹/1,000 picks (current)": g.at["baseline", "labour_cost_per_1000_picks"],
                         "₹/1,000 picks (optimized)": g.at["optimized", "labour_cost_per_1000_picks"], "Labour saving": None,
                         "Heavy violations (current)": int(g.at["baseline", "heavy_slot_violations"]),
                         "Heavy violations (optimized)": int(g.at["optimized", "heavy_slot_violations"]),
                         "Bound captured": None})
        t = pd.DataFrame(rows)
        section(ws, 4, 2, "Current vs optimized vs perfect-information bound (formulas)", 16)
        r0, r1, c0, c1 = write_df(ws, t, 5, 2, {"Pick lines": "int", "Travel m/line (current)": "dec2",
                                                 "Travel m/line (optimized)": "dec2", "Travel m/line (bound)": "dec2",
                                                 "Travel reduction": "pct", "Picks/h (current)": "dec1", "Picks/h (optimized)": "dec1",
                                                 "Picks/h (bound)": "dec1", "Throughput gain": "pct", "₹/1,000 picks (current)": "inr",
                                                 "₹/1,000 picks (optimized)": "inr", "Labour saving": "pct",
                                                 "Bound captured": "pct"}, "WarehousePerf")
        for r in range(r0 + 1, r1 + 1):
            ws.cell(r, c0 + 5, f"=1-{L(c0 + 3)}{r}/{L(c0 + 2)}{r}").number_format = FMT["pct"]
            ws.cell(r, c0 + 9, f"={L(c0 + 7)}{r}/{L(c0 + 6)}{r}-1").number_format = FMT["pct"]
            ws.cell(r, c0 + 12, f"=1-{L(c0 + 11)}{r}/{L(c0 + 10)}{r}").number_format = FMT["pct"]
            ws.cell(r, c0 + 15, f"=({L(c0 + 2)}{r}-{L(c0 + 3)}{r})/({L(c0 + 2)}{r}-{L(c0 + 4)}{r})").number_format = FMT["pct"]
        tr = r1 + 1
        ws.cell(tr, c0, "All cities (line-weighted)").font = font(10, True)
        for k, kind in ((5, "travel"), (9, "thr")):
            f = (f"=1-SUMPRODUCT({L(c0 + 3)}{r0 + 1}:{L(c0 + 3)}{r1},$C${r0 + 1}:$C${r1})/SUMPRODUCT({L(c0 + 2)}{r0 + 1}:"
                 f"{L(c0 + 2)}{r1},$C${r0 + 1}:$C${r1})") if kind == "travel" else (
                f"=SUMPRODUCT({L(c0 + 7)}{r0 + 1}:{L(c0 + 7)}{r1},$C${r0 + 1}:$C${r1})/SUMPRODUCT({L(c0 + 6)}{r0 + 1}:"
                f"{L(c0 + 6)}{r1},$C${r0 + 1}:$C${r1})-1")
            c = ws.cell(tr, c0 + k, f)
            c.number_format = FMT["pct"]
            c.font = font(10, True)
        self.ref["wh_travel"] = f"{q(s)}!{L(c0 + 5)}{tr}"
        self.ref["wh_thr"] = f"{q(s)}!{L(c0 + 9)}{tr}"
        self.ref["wh_tbl"] = (s, r0, r1, c0)
        data_bar(ws, col_range(r0 + 1, r1, c0 + 5), TEAL)
        data_bar(ws, col_range(r0 + 1, r1, c0 + 9), LEAF)
        data_bar(ws, col_range(r0 + 1, r1, c0 + 15), MARIGOLD)
        bar_chart(ws, f"B{tr + 2}", ws, (c0, r0 + 1, r1), [c0 + 6, c0 + 7, c0 + 8], r0, "Picks per hour: current, optimized, bound",
                  "picks/h", colors=[SLATE, TEAL, MARIGOLD], w=16, h=8)
        bar_chart(ws, f"J{tr + 2}", ws, (c0, r0 + 1, r1), [c0 + 2, c0 + 3, c0 + 4], r0, "Walking metres per pick line",
                  "m/line", colors=[SLATE, TEAL, MARIGOLD], w=16, h=8)
        note(ws, tr + 19, 2, "Targets: 35% less picker travel and 25-30% more throughput. Measured gains sit within ~1-2 points of "
             "the perfect-information bound (same method fed each SKU's long-run demand rate), so the shortfall is a property of "
             "this SKU profile and return-routing batch picking, not of the slotting method. Reaching the targets would need a "
             "different picking policy (e.g. zone or wave picking with S-shape routing) - outside slotting scope.", 16)
        row = tr + 22
        zs = self.p("warehouse_zone_summary.parquet").pivot_table(index="city_id", columns="opt_zone", values="pick_share")
        zs = zs.reindex(self.cities).reset_index()
        zs.insert(0, "City", zs["city_id"].map(self.names))
        zs = zs.drop(columns="city_id").rename(columns={"A": "Zone A (golden)", "B": "Zone B", "C": "Zone C (bulk)"})
        section(ws, row, 2, "Share of training picks by zone after slotting", 6,
                "Zone A holds 20% of slots nearest dispatch; velocity zoning concentrates picks there.")
        r0, r1, c0, c1 = write_df(ws, zs, row + 2, 2, {c: "pct" for c in zs.columns[1:]}, "ZoneShare")
        bar_chart(ws, f"B{r1 + 2}", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1 + 1)), r0, "Pick share by zone",
                  "share", stacked=True, colors=[TEAL, MARIGOLD, SLATE], w=16, h=7, number_format="0%")
        cv = self.p("warehouse_shrinkage_cv.parquet")
        cvp = cv.pivot_table(index="shrinkage_k", columns="city_id", values="validation_travel_m")[self.cities].reset_index()
        cvp.columns = ["Shrinkage k"] + [self.names[c] for c in self.cities]
        section(ws, row, 10, "Velocity shrinkage chosen by internal validation (travel m)", 7)
        r0, r1, c0, _ = write_df(ws, cvp, row + 2, 10, {c: "int" for c in cvp.columns[1:]}, "ShrinkageCV")
        for c in range(c0 + 1, c0 + len(cvp.columns)):
            color_scale(ws, col_range(r0 + 1, r1, c), reverse=True)

    # ------------------------------------------------------------------ 9
    def simulation(self) -> None:
        s = SHEETS[8]
        ws = self.ws[s]
        sm = self.p("simulation_summary.parquet")
        reps = int(sm["replications"].iloc[0])
        setup_sheet(ws, "Simulation Results", f"SimPy DC model per city · {reps} Monte Carlo replications per scenario · "
                    f"target P95 intake-to-dispatch <= {self.cfg['simulation']['cycle_time_target_min']} min", 17)
        sm = sm.copy()
        sm["City"] = sm["city_id"].map(self.names)
        sm["Roster (pick/pack/QC/load)"] = (sm["staff_pickers"].astype(str) + "/" + sm["staff_packers"].astype(str) + "/"
                                            + sm["staff_qc"].astype(str) + "/" + sm["staff_loaders"].astype(str))
        piv = sm.pivot_table(index="City", columns="scenario", values="pooled_cycle_p95").reindex([self.names[c] for c in self.cities])
        scen_order = ["Baseline staffing", "Optimized staffing", "Diwali peak - baseline staffing", "Diwali peak - peak roster"]
        piv = piv[scen_order]
        piv["Target"] = self.cfg["simulation"]["cycle_time_target_min"]
        piv = piv.reset_index()
        section(ws, 4, 2, "Pooled P95 cycle time (minutes) by scenario", 7)
        r0, r1, c0, c1 = write_df(ws, piv, 5, 2, {c: "dec1" for c in piv.columns[1:]}, "SimP95")
        color_scale(ws, f"{L(c0 + 1)}{r0 + 1}:{L(c1 - 1)}{r1}", reverse=True)
        bar = BarChart()
        for k, col in enumerate(range(c0 + 1, c1)):
            bar.add_data(Reference(ws, min_col=col, min_row=r0, max_row=r1), titles_from_data=True)
            bar.series[-1].graphicalProperties.solidFill = [SLATE, TEAL, SIGNAL, MARIGOLD][k]
        bar.set_categories(Reference(ws, min_col=c0, min_row=r0 + 1, max_row=r1))
        ln = LineChart()
        ln.add_data(Reference(ws, min_col=c1, min_row=r0, max_row=r1), titles_from_data=True)
        ln.series[0].graphicalProperties.line.solidFill = INK
        ln.series[0].graphicalProperties.line.dashStyle = "dash"
        bar.y_axis.title = "minutes"
        bar.title = "P95 cycle time vs 2-hour target"
        bar.width, bar.height = 18, 9
        bar.legend.position = "b"
        bar.x_axis.delete = False
        bar.y_axis.delete = False
        bar += ln
        ws.add_chart(bar, "J4")
        self.ref["sim_tbl"] = (s, r0, r1, c0)
        cols = ["City", "scenario", "Roster (pick/pack/QC/load)", "daily_staff_cost", "throughput_p10", "throughput_p50",
                "throughput_p90", "cycle_p95_p10", "cycle_p95_p50", "cycle_p95_p90", "pooled_cycle_p95", "sla_compliance_p50",
                "backlog_p50", "util_pickers", "util_packers", "util_loaders", "bottleneck", "meets_p95_target"]
        det = sm[cols].copy()
        det.columns = ["City", "Scenario", "Roster", "Staff ₹/day", "Throughput P10", "Throughput P50", "Throughput P90",
                       "P95 cycle P10", "P95 cycle P50", "P95 cycle P90", "Pooled P95", "SLA share P50", "Backlog P50",
                       "Picker util", "Packer util", "Loader util", "Bottleneck", "Meets target"]
        dr = r1 + 14
        section(ws, dr, 2, "Scenario detail (P10 / P50 / P90 across replications)", 18)
        r0, r1, c0, c1 = write_df(ws, det, dr + 1, 2, {"Staff ₹/day": "inr", "Throughput P10": "int", "Throughput P50": "int",
                                                        "Throughput P90": "int", "P95 cycle P10": "dec1", "P95 cycle P50": "dec1",
                                                        "P95 cycle P90": "dec1", "Pooled P95": "dec1", "SLA share P50": "pct",
                                                        "Backlog P50": "dec1", "Picker util": "pct", "Packer util": "pct",
                                                        "Loader util": "pct"}, "SimDetail", {"Scenario": 28})
        for k in (13, 14, 15):
            color_scale(ws, col_range(r0 + 1, r1, c0 + k), reverse=True)
        ws.conditional_formatting.add(col_range(r0 + 1, r1, c0 + 17), CellIsRule(operator="equal", formula=['"Yes"'],
                                      fill=PatternFill("solid", fgColor="D7EEE0")))
        ws.conditional_formatting.add(col_range(r0 + 1, r1, c0 + 17), CellIsRule(operator="equal", formula=['"No"'],
                                      fill=PatternFill("solid", fgColor="F8D9D6")))
        h = self.p("simulation_cycle_hist.parquet")
        h = h[h["city_id"] == "MUM"].pivot_table(index="bin_start_min", columns="scenario", values="orders")[scen_order].reset_index()
        h.columns = ["Cycle bin (min)"] + scen_order
        hr = r1 + 3
        section(ws, hr, 2, "Mumbai DC: distribution of intake-to-dispatch time (all replications)", 6)
        r0, r1, c0, c1 = write_df(ws, h, hr + 1, 2, {c: "int" for c in h.columns[1:]}, "SimHist")
        line_chart(ws, f"H{hr}", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1 + 1)), r0,
                   "Mumbai: cycle-time distribution by scenario", "orders", "minutes (bin start)",
                   colors=[SLATE, TEAL, SIGNAL, MARIGOLD], w=18, h=9)
        ws.freeze_panes = "C4"

    # ------------------------------------------------------------------ 10
    def network(self) -> None:
        s = SHEETS[9]
        ws = self.ws[s]
        setup_sheet(ws, "Network Design", "p-median (PuLP) over candidate MFC sites per city · 2-hour SLA tested at evening-peak traffic", 17)
        rk = self.p("network_city_ranking.parquet").copy()
        rk["best_npv"] = rk["best_npv"] / 1e7
        rk = rk[["investment_priority", "city", "state", "parcels_per_day", "baseline_sla_2h", "best_p", "best_sites", "best_npv",
                 "break_even_parcels_per_day", "break_even_multiplier", "min_p_full_sla", "cost_per_order_km", "recommendation"]]
        rk.columns = ["Priority", "City", "State(s)", "Parcels/day", "2h SLA today", "Best p", "Best sites", "Best NPV (₹ cr)",
                      "Break-even parcels/day", "Break-even x volume", "Min p for 100% SLA", "₹ / order-km", "Recommendation"]
        section(ws, 4, 2, "Cross-state MFC investment ranking (lowest break-even volume first)", 14)
        r0, r1, c0, c1 = write_df(ws, rk, 5, 2, {"Parcels/day": "dec1", "2h SLA today": "pct", "Best NPV (₹ cr)": "cr",
                                                  "Break-even parcels/day": "int", "Break-even x volume": "x", "₹ / order-km": "inr2"},
                                  "CityRanking", {"Best sites": 22, "Recommendation": 70, "State(s)": 22})
        data_bar(ws, col_range(r0 + 1, r1, c0 + 9), MARIGOLD)
        color_scale(ws, col_range(r0 + 1, r1, c0 + 4))
        self.ref["net_rank"] = (s, r0, r1, c0)
        sc = self.p("network_scenarios.parquet")
        sc["City"] = sc["city_id"].map(self.names)
        names = [self.names[c] for c in self.cities]
        row = r1 + 3
        for val, title, fmt, ylab in (("avg_km_per_order", "Average last-mile km per parcel vs number of new MFCs", "dec2", "km"),
                                      ("sla_2h_coverage", "2-hour SLA coverage vs number of new MFCs", "pct", "share"),
                                      ("npv", "Five-year NPV vs number of new MFCs (₹)", "inr", "₹")):
            pv = sc.pivot_table(index="p", columns="City", values=val)[names].reset_index().rename(columns={"p": "New MFCs (p)"})
            section(ws, row, 2, title, 7)
            r0p, r1p, c0p, c1p = write_df(ws, pv, row + 1, 2, {c: fmt for c in pv.columns[1:]}, None)
            if val == "npv":
                bar_chart(ws, f"J{row}", ws, (c0p, r0p + 2, r1p), list(range(c0p + 1, c1p + 1)), r0p, title, ylab,
                          x_title="p", w=17, h=8, number_format='"₹"#,##0,,"M"')
            else:
                line_chart(ws, f"J{row}", ws, (c0p, r0p + 1, r1p), list(range(c0p + 1, c1p + 1)), r0p, title, ylab, "p",
                           markers=True, w=17, h=8, number_format="0%" if fmt == "pct" else None)
            row = max(r1p, row + 15) + 2
        det = sc[["City", "p", "site_names", "avg_km_per_order", "sla_2h_coverage", "mfc_orders_per_day", "capex", "annual_opex",
                  "annual_linehaul", "annual_saving_y1", "npv", "break_even_parcels_per_day", "solver_status"]].copy()
        det.columns = ["City", "p", "Sites", "Km/parcel", "2h SLA", "MFC parcels/day", "Capex ₹", "Opex ₹/yr", "Line-haul ₹/yr",
                       "Saving Y1 ₹", "NPV ₹", "Break-even parcels/day", "Solver"]
        section(ws, row, 2, "All configurations", 13)
        r0, r1, c0, c1 = write_df(ws, det, row + 1, 2, {"Km/parcel": "dec2", "2h SLA": "pct", "MFC parcels/day": "dec1",
                                                         "Capex ₹": "inr", "Opex ₹/yr": "inr", "Line-haul ₹/yr": "inr",
                                                         "Saving Y1 ₹": "inr", "NPV ₹": "inr", "Break-even parcels/day": "int"},
                                  "NetworkScenarios", {"Sites": 40})
        color_scale(ws, col_range(r0 + 1, r1, c0 + 10))
        sites = self.p("network_sites.parquet")[["site_id", "city", "name", "rent_psf_month", "annual_fixed_opex_inr",
                                                 "linehaul_road_km", "selected_in_p"]]
        sites.columns = ["Site", "City", "Location", "Rent ₹/sqft/month", "Fixed opex ₹/yr", "Line-haul km from DC", "Selected in p"]
        sr = r1 + 3
        section(ws, sr, 2, "Candidate MFC sites", 8)
        r0, r1, c0, _ = write_df(ws, sites, sr + 1, 2, {"Fixed opex ₹/yr": "inr", "Line-haul km from DC": "dec1"}, "Sites",
                                 {"Location": 20})
        data_bar(ws, col_range(r0 + 1, r1, c0 + 3), SIGNAL)
        ws.freeze_panes = "C6"

    # ------------------------------------------------------------------ 11
    def economics(self) -> None:
        s = SHEETS[10]
        ws = self.ws[s]
        setup_sheet(ws, "Facility Economics", "Live five-year NPV of each city's best MFC configuration · edit WACC/growth on "
                    "'Master Data & Assumptions' · Python tie-out column", 18)
        ws.cell(4, 2, "WACC").font = font(10, True)
        c = ws.cell(4, 3, f"={self.ref['WACC']}")
        c.font = font(10, True, GREEN)
        c.number_format = FMT["pct"]
        ws.cell(5, 2, "Demand growth").font = font(10, True)
        c = ws.cell(5, 3, f"={self.ref['Demand growth / year']}")
        c.font = font(10, True, GREEN)
        c.number_format = FMT["pct"]
        ws.cell(4, 5, "Year t").font = font(10, True)
        for t in range(1, 6):
            ws.cell(4, 5 + t, t).font = font(10, True)
        sc = self.p("network_scenarios.parquet")
        rk = self.p("network_city_ranking.parquet").set_index("city_id")
        hdr = ["City", "p", "Capex ₹", "Fixed opex ₹/yr", "Line-haul ₹/yr", "Saving Y1 ₹", "CF Y0", "CF Y1", "CF Y2", "CF Y3",
               "CF Y4", "CF Y5", "NPV (formula)", "NPV (Python)", "Tie-out diff", "Break-even x (formula)", "Payback"]
        hr = 7
        section(ws, hr - 1, 2, "Cash flows and NPV (₹)", 17)
        for k, h in enumerate(hdr):
            cell = ws.cell(hr, 2 + k, h)
            cell.font = font(10, True, "FFFFFF")
            cell.fill = PatternFill("solid", fgColor=INK)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[hr].height = 30
        yrs = "$F$4:$J$4"
        for i, cid in enumerate(self.cities):
            r = hr + 1 + i
            p = int(rk.at[cid, "best_p"])
            row = sc[(sc["city_id"] == cid) & (sc["p"] == p)].iloc[0]
            vals = [self.names[cid], p, row["capex"], row["annual_opex"], row["annual_linehaul"], row["annual_saving_y1"]]
            for k, v in enumerate(vals):
                cell = ws.cell(r, 2 + k, float(v) if k > 1 else v)
                cell.font = font(10)
                if k > 1:
                    cell.number_format = FMT["inr"]
            ws.cell(r, 8, f"=-D{r}").number_format = FMT["inr"]
            for t in range(1, 6):
                ws.cell(r, 8 + t, f"=$G{r}*(1+$C$5)^({t}-1)-$E{r}-$F{r}").number_format = FMT["inr"]
            ws.cell(r, 14, f"=H{r}+SUMPRODUCT(I{r}:M{r},1/(1+$C$4)^{yrs})").number_format = FMT["inr"]
            ws.cell(r, 15, float(row["npv"])).number_format = FMT["inr"]
            ws.cell(r, 16, f"=N{r}-O{r}").number_format = FMT["inr"]
            ws.cell(r, 17, f"=IFERROR((D{r}+(E{r}+F{r})*SUMPRODUCT(1/(1+$C$4)^{yrs}))/(G{r}*SUMPRODUCT((1+$C$5)^({yrs}-1)"
                           f"/(1+$C$4)^{yrs})),\"n/a\")").number_format = FMT["x"]
            ws.cell(r, 18, f"=IF(SUM(H{r}:M{r})<0,\"> 5 yrs\",\"within 5 yrs\")")
        r_last = hr + len(self.cities)
        tot = r_last + 1
        ws.cell(tot, 2, "All cities").font = font(10, True)
        for col in range(4, 15):
            cell = ws.cell(tot, col, f"=SUM({L(col)}{hr + 1}:{L(col)}{r_last})")
            cell.number_format = FMT["inr"]
            cell.font = font(10, True)
        self.ref["npv_total"] = f"{q(s)}!N{tot}"
        self.ref["econ_rows"] = (s, hr, r_last)
        color_scale(ws, f"N{hr + 1}:N{r_last}")
        data_bar(ws, f"Q{hr + 1}:Q{r_last}", MARIGOLD)
        ws.conditional_formatting.add(f"P{hr + 1}:P{r_last}", FormulaRule(formula=[f"ABS(P{hr + 1})>1000"],
                                      fill=PatternFill("solid", fgColor="F8D9D6")))
        ch = BarChart()
        for col in range(8, 14):
            ch.add_data(Reference(ws, min_col=col, min_row=hr, max_row=r_last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, min_row=hr + 1, max_row=r_last))
        for k, srs in enumerate(ch.series):
            srs.graphicalProperties.solidFill = SIGNAL if k == 0 else [TEAL, "3C8D5A", MARIGOLD, "8E5A9B", SLATE][k - 1]
        ch.title = "Cash flow by year for each city's best configuration (₹)"
        ch.y_axis.numFmt = '"₹"#,##0,,"M"'
        ch.width, ch.height = 20, 9
        ch.legend.position = "b"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ws.add_chart(ch, f"B{tot + 2}")
        # Sensitivity
        sen = self.p("network_sensitivity.parquet")
        sen["label"] = sen["lever"] + " " + sen["value"].map(lambda v: f"{v:+.0%}" if abs(v) < 1 else f"{v}")
        sen.loc[sen["lever"].isin(["wacc", "growth"]), "label"] = sen["lever"] + " " + sen["value"].map(lambda v: f"{v:.0%}")
        piv = sen.pivot_table(index="label", columns="city_id", values="npv")[self.cities] / 1e7
        piv.columns = [self.names[c] for c in self.cities]
        piv = piv.reset_index().rename(columns={"label": "Lever setting"})
        sr = tot + 22
        section(ws, sr, 2, "One-way sensitivity of best-configuration NPV (₹ crore)", 8,
                "Levers: WACC, MFC capex, parcel growth, variable last-mile cost per order-km.")
        r0, r1, c0, c1 = write_df(ws, piv, sr + 2, 2, {c: "cr" for c in piv.columns[1:]}, "NPVSensitivity", {"Lever setting": 22})
        color_scale(ws, f"{L(c0 + 1)}{r0 + 1}:{L(c1)}{r1}")
        bar_chart(ws, f"J{sr}", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1 + 1)), r0, "NPV sensitivity by lever (₹ crore)",
                  "₹ crore", horizontal=True, w=18, h=10)
        note(ws, r1 + 2, 2, "Tie-out: 'NPV (formula)' is recomputed in Excel from the cash-flow inputs; 'NPV (Python)' is the "
             "network_designer output at base assumptions. A non-zero difference appears only when the WACC or growth levers "
             "are edited. Break-even = volume multiple at which NPV = 0 with the same sites.", 16)

    # ------------------------------------------------------------------ 12
    def ml(self) -> None:
        s = SHEETS[11]
        ws = self.ws[s]
        setup_sheet(ws, "ML Performance", "Time-ordered validation (last 21 days unseen) · models used only where they beat a simple baseline", 17)
        perf = self.p("ml_performance.parquet")
        tt = perf[perf["model"] == "Travel time"].set_index("approach")
        dm = perf[perf["model"] == "Zone daily demand"].set_index("approach")
        tb, tm = tt.iloc[0], tt.iloc[1]
        db, dmm = dm.iloc[0], dm.iloc[1]
        kpi_card(ws, 4, 2, "Travel-time MAE reduction vs planning formula", 1 - tm["MAE"] / tb["MAE"], FMT["pct"], 4,
                 sub=f"{tb['MAE']:.1f} -> {tm['MAE']:.1f} min")
        kpi_card(ws, 4, 6, "Demand MAE reduction vs seasonal naive", 1 - dmm["MAE"] / db["MAE"], FMT["pct"], 4,
                 sub=f"{db['MAE']:.2f} -> {dmm['MAE']:.2f} orders/zone-day")
        kpi_card(ws, 4, 10, "Validation window starts", str(perf["validation_start"].iloc[0]), "@", 4, SLATE)
        section(ws, 7, 2, "Validation metrics", 9)
        r0, r1, c0, c1 = write_df(ws, perf[["model", "approach", "MAE", "RMSE", "MAPE", "n", "unit", "validation_start"]],
                                  8, 2, {"MAE": "dec2", "RMSE": "dec2", "MAPE": '0.0"%"', "n": "int"}, "MLMetrics",
                                  {"approach": 30, "model": 18})
        bar_chart(ws, "L7", ws, (c0 + 1, r0 + 1, r1), [c0 + 2, c0 + 3], r0, "Validation error: baselines vs models", "error",
                  colors=[TEAL, MARIGOLD], horizontal=True, w=16, h=7)
        zf = self.p("ml_zone_factors.parquet")
        zcity = self.b.addresses.drop_duplicates("zone").set_index("zone")["city_id"]
        zf["City"] = zf["dest_zone"].map(zcity).map(self.names)
        zf = zf.sort_values(["City", "dest_zone"])[["City", "dest_zone", "factor"]]
        zf.columns = ["City", "Zone", "Travel-time factor vs plan"]
        zr = r1 + 3
        section(ws, zr, 2, "Learned zone travel-time calibration (fed to routing and fleet sizing)", 6,
                "Factor > 1 = slower than the network-wide planning formula (congestion, river crossings); < 1 = faster.")
        r0, r1, c0, _ = write_df(ws, zf, zr + 2, 2, {"Travel-time factor vs plan": "dec3"}, "ZoneFactors", {"Zone": 20})
        color_scale(ws, col_range(r0 + 1, r1, c0 + 2), reverse=True)
        bar_chart(ws, f"G{zr}", ws, (c0 + 1, r0 + 1, r1), [c0 + 2], r0, "Zone travel-time factor", "x plan", horizontal=True,
                  colors=[TEAL], w=13, h=14)
        fi = self.p("ml_feature_importance.parquet").head(12)
        section(ws, zr, 15, "Top travel-time features", 4)
        r0f, r1f, c0f, _ = write_df(ws, fi, zr + 2, 15, {"importance": "dec3"}, "FeatureImportance", {"feature": 24})
        data_bar(ws, col_range(r0f + 1, r1f, c0f + 1), TEAL)
        dv = self.p("ml_demand_validation.parquet").groupby("date")[["orders", "ml_forecast", "seasonal_naive"]].sum().reset_index()
        dv.columns = ["Date", "Actual", "Gradient boosting", "Seasonal naive"]
        vr = r1 + 3
        section(ws, vr, 2, "Demand validation: all-city daily orders (unseen 21 days)", 5)
        r0, r1, c0, c1 = write_df(ws, dv, vr + 1, 2, {"Actual": "int", "Gradient boosting": "dec1", "Seasonal naive": "dec1"}, "DemandVal")
        line_chart(ws, f"G{vr}", ws, (c0, r0 + 1, r1), [c0 + 1, c0 + 2, c0 + 3], r0, "Actual vs forecasts (validation)", "orders",
                   colors=[INK, TEAL, MARIGOLD], dashed={2}, markers=True, w=18, h=8)
        fc = self.p("ml_demand_forecast.parquet").groupby(["date", "city_id"])["forecast"].sum().unstack()[self.cities]
        fc.columns = [self.names[c] for c in self.cities]
        fc = fc.reset_index().rename(columns={"date": "Date"})
        fr = r1 + 3
        section(ws, fr, 2, "Next 14 days: forecast orders by city (feeds workload and fleet scenario selection)", 7)
        r0, r1, c0, c1 = write_df(ws, fc, fr + 1, 2, {c: "dec1" for c in fc.columns[1:]}, "Forecast14")
        line_chart(ws, f"J{fr}", ws, (c0, r0 + 1, r1), list(range(c0 + 1, c1 + 1)), r0, "14-day forecast by city", "orders/day",
                   markers=True, w=17, h=8)

    # ------------------------------------------------------------------ 13
    def kpis(self) -> None:
        s = SHEETS[12]
        ws = self.ws[s]
        setup_sheet(ws, "KPI Definitions", "How every reported metric is calculated", 8)
        rows = [
            ("Last-mile cost", "Fleet cost of all routes plus spot-outsourcing cost of unserved orders", "Σ(fixed + km x ₹/km + duty h x wage) + Σ outsourcing", "₹/day", "Routing", "15-22% reduction"),
            ("Fleet operating cost", "Cost of the routes the fleet drives", "Σ(fixed + km x ₹/km + duty h x wage)", "₹/day", "Routing", "-"),
            ("Time-window compliance", "Orders delivered inside their window / all orders (unserved = non-compliant)", "on-time / orders", "%", "Routing, Fleet", "> 95%"),
            ("Straight-line km", "Great-circle (Haversine) distance between stops", "Haversine", "km", "Routing", "-"),
            ("Road km (estimate)", "Straight-line km x planning circuity 1.35 (no street routing)", "haversine x 1.35", "km", "Routing", "-"),
            ("Duty overrun", "Route duration replayed with departure-time traffic exceeds the duty limit", "duration > limit", "count", "Routing", "0"),
            ("Recommended fleet", "Vehicles used by the cost-minimising fleet-size-and-mix solution from a candidate pool", "count of used vehicles", "vehicles", "Fleet", "min feasible"),
            ("Annual TCO (own)", "Equivalent annual capital cost + insurance + maintenance + fuel + labour", "EAC(capex, residual, WACC, life) + opex", "₹/yr", "Fleet", "-"),
            ("Picker travel", "Walking metres per pick line, held-out orders, return routing, batches of 3", "Σ tour m / lines", "m/line", "Warehouse", "-35%"),
            ("Picks per hour", "Pick lines / (walking + handling + tour setup hours)", "lines / hours", "lines/h", "Warehouse", "+25-30%"),
            ("Perfect-information bound", "Same slotting method fed each SKU's long-run demand rate", "upper bound", "%", "Warehouse", "-"),
            ("Pooled P95 cycle time", "95th percentile of intake-to-dispatch minutes across all simulated orders", "quantile 0.95", "min", "Simulation", "<= 120"),
            ("P10/P50/P90", "Percentiles across Monte Carlo replications", "quantiles", "-", "Simulation", "-"),
            ("2-hour SLA coverage", "Share of parcel demand within 120 min (40 min processing + travel at 2.2x evening traffic)", "Σ w(i) [t(i) <= 80] / Σ w", "%", "Network", "-"),
            ("NPV (5-year)", "-capex + Σ (saving(t) - opex - line-haul) / (1+WACC)^t, no terminal value", "NPV", "₹", "Network", "quantified"),
            ("Break-even volume", "Multiple of today's parcel volume at which NPV = 0 with the same sites", "PV(costs)/PV(savings)", "x", "Network", "-"),
            ("Travel-time MAE", "Mean absolute error of leg minutes on the unseen validation window", "mean |pred - actual|", "min", "ML", "beat formula"),
            ("Logistics health score", "Delivery 30 (90%->99%, less duty overruns), cost 20 (/22%), throughput 15 (/30%), DC readiness 20 (optimized + Diwali-today), fleet scalability 15 (current fleet)", "Σ weight x scaled component", "0-100", "Command Center", "-"),
        ]
        df = pd.DataFrame(rows, columns=["KPI", "Definition", "Formula", "Unit", "Module", "Target"])
        write_df(ws, df, 4, 2, None, "KPIDefs", {"KPI": 26, "Definition": 70, "Formula": 40, "Module": 16, "Target": 16})

    # ------------------------------------------------------------------ 14
    def comparison(self) -> None:
        s = SHEETS[13]
        ws = self.ws[s]
        setup_sheet(ws, "Scenario Comparison", "Business targets vs measured results (live links) · cross-module view by city", 12)
        hdr = ["Target", "Low", "High", "Actual", "Status", "Explanation"]
        for k, h in enumerate(hdr):
            c = ws.cell(4, 2 + k, h)
            c.font = font(10, True, "FFFFFF")
            c.fill = PatternFill("solid", fgColor=INK)
        targets = [
            ("Last-mile cost reduction (incl. outsourcing)", 0.15, 0.22, f"={self.ref['routing_lm_red']}",
             "Measured on the representative day against a window-aware manual baseline, same evaluator."),
            ("Fleet operating cost reduction", 0.15, 0.22, f"={self.ref['routing_fleet_red']}", "Routes only, excludes outsourcing."),
            ("Delivery time-window compliance", 0.95, 1.0, f"={self.ref['routing_comp']}", "Unserved orders counted as non-compliant."),
            ("Warehouse throughput gain", 0.25, 0.30, f"={self.ref['wh_thr']}",
             "Within ~1-2 pts of the perfect-information bound; picking policy limits further gains."),
            ("Picker travel reduction", 0.35, 0.35, f"={self.ref['wh_travel']}", "Bound ~27%: target not reachable with return routing."),
        ]
        for i, (name, lo, hi, f, why) in enumerate(targets):
            r = 5 + i
            ws.cell(r, 2, name).font = font(10)
            for k, v in ((3, lo), (4, hi)):
                c = ws.cell(r, k, v)
                c.number_format = FMT["pct"]
                c.font = font(10, color="0000FF")
            c = ws.cell(r, 5, f)
            c.number_format = FMT["pct"]
            c.font = font(10, True, GREEN)
            ws.cell(r, 6, f'=IF(E{r}>=C{r},IF(E{r}>D{r},"Exceeds","Met"),"Below target")').font = font(10, True)
            ws.cell(r, 7, why).font = font(9, color=SLATE)
        r = 5 + len(targets)
        ws.cell(r, 2, "Five-year facility NPV (sum of best configs)").font = font(10)
        c = ws.cell(r, 5, f"={self.ref['npv_total']}")
        c.number_format = FMT["inr"]
        c.font = font(10, True, GREEN)
        ws.cell(r, 6, "Quantified").font = font(10, True)
        ws.cell(r, 7, "Negative at today's volumes in every city; break-even volumes on Network Design.").font = font(9, color=SLATE)
        rng = f"F5:F{r}"
        for text, fill in (("Exceeds", "D7EEE0"), ("Met", "E9F5EC"), ("Below target", "F8D9D6"), ("Quantified", "FFF6DD")):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'F5="{text}"'], fill=PatternFill("solid", fgColor=fill)))
        ws.column_dimensions["B"].width = 44
        ws.column_dimensions["G"].width = 80
        # Cross-module city matrix
        cm = self.p("routing_city_summary.parquet")
        b = cm[cm["plan"] == "baseline"].set_index("city_id")
        o = cm[cm["plan"] == "optimized"].set_index("city_id")
        whs = self.j("warehouse_summary.json")["by_city"]
        sm = self.p("simulation_summary.parquet").set_index(["city_id", "scenario"])
        rk = self.p("network_city_ranking.parquet").set_index("city_id")
        fs = self.p("fleet_scenarios.parquet")
        rows = []
        for cid in self.cities:
            f1 = fs[(fs["city_id"] == cid) & (fs["demand_multiplier"] == 1.0)].iloc[0]
            rows.append({"City": self.names[cid], "Last-mile cost reduction": 1 - o.at[cid, "last_mile_cost"] / b.at[cid, "last_mile_cost"],
                         "Optimized compliance": o.at[cid, "tw_compliance"], "Fleet at 1.0x": int(f1["recommended_fleet"]),
                         "Pick travel reduction": whs[cid]["comparison"]["travel_reduction_pct"] / 100,
                         "Throughput gain": whs[cid]["comparison"]["throughput_gain_pct"] / 100,
                         "P95 now (min)": sm.at[(cid, "Baseline staffing"), "pooled_cycle_p95"],
                         "P95 optimized (min)": sm.at[(cid, "Optimized staffing"), "pooled_cycle_p95"],
                         "Diwali P95 current roster": sm.at[(cid, "Diwali peak - baseline staffing"), "pooled_cycle_p95"],
                         "MFC break-even x": rk.at[cid, "break_even_multiplier"],
                         "2h SLA today": rk.at[cid, "baseline_sla_2h"]})
        mt = pd.DataFrame(rows)
        mr = r + 3
        section(ws, mr, 2, "City scorecard across modules", 11)
        r0, r1, c0, c1 = write_df(ws, mt, mr + 1, 2, {"Last-mile cost reduction": "pct", "Optimized compliance": "pct",
                                                       "Fleet at 1.0x": "int", "Pick travel reduction": "pct", "Throughput gain": "pct",
                                                       "P95 now (min)": "dec1", "P95 optimized (min)": "dec1",
                                                       "Diwali P95 current roster": "dec1", "MFC break-even x": "x",
                                                       "2h SLA today": "pct"}, "CityScorecard")
        for k in (1, 2, 4, 5, 10):
            color_scale(ws, col_range(r0 + 1, r1, c0 + k))
        for k in (6, 7, 8, 9):
            color_scale(ws, col_range(r0 + 1, r1, c0 + k), reverse=True)
        self.ref["scorecard"] = (s, r0, r1, c0)

    # ------------------------------------------------------------------ 15
    def executive(self) -> None:
        s = SHEETS[14]
        ws = self.ws[s]
        cc = self.j("command_center_view.json")
        setup_sheet(ws, "Executive Dashboard", "Logistics and Network Design · Command Center · Mumbai, Delhi NCR, Bengaluru, Chennai, Kolkata", 17, MARIGOLD)
        kpi_card(ws, 4, 2, "Logistics health score (0-100)", cc["health"]["score"], "0.0", 2, INK)
        kpi_card(ws, 4, 4, "Last-mile cost reduction", f"={self.ref['routing_lm_red']}", FMT["pct"], 2)
        kpi_card(ws, 4, 6, "In-window delivery", f"={self.ref['routing_comp']}", FMT["pct"], 2, LEAF)
        kpi_card(ws, 4, 8, "Vehicles saved per day", f"={self.ref['veh_base']}-{self.ref['veh_opt']}", FMT["int"], 2)
        kpi_card(ws, 4, 10, "Picker travel reduction", f"={self.ref['wh_travel']}", FMT["pct"], 2)
        kpi_card(ws, 4, 12, "Pick throughput gain", f"={self.ref['wh_thr']}", FMT["pct"], 2, LEAF)
        kpi_card(ws, 4, 14, "5-yr NPV, best MFC configs", f"={self.ref['npv_total']}/10^7", FMT["cr"], 2, SIGNAL)
        rk = self.p("network_city_ranking.parquet")
        kpi_card(ws, 4, 16, "First city for MFC investment", rk.iloc[0]["city"], "@", 2, MARIGOLD,
                 sub=f"break-even {rk.iloc[0]['break_even_multiplier']:.1f}x volume")
        rs, r0, r1, c0 = self.ref["routing_tbl"]
        bar_chart(ws, "B8", self.ws[rs], (c0, r0 + 1, r1), [c0 + 2, c0 + 3], r0, "Daily last-mile cost: baseline vs optimized",
                  "₹/day", colors=[SLATE, TEAL], w=15, h=7.5, number_format='"₹"#,##0,"K"')
        ss, s0, s1, sc0 = self.ref["sim_tbl"]
        bar_chart(ws, "J8", self.ws[ss], (sc0, s0 + 1, s1), [sc0 + 1, sc0 + 2, sc0 + 3, sc0 + 4], s0,
                  "DC P95 cycle time (min) by roster scenario", "minutes", colors=[SLATE, TEAL, SIGNAL, MARIGOLD], w=16, h=7.5)
        ws_, w0, w1, wc0 = self.ref["wh_tbl"]
        bar_chart(ws, "B23", self.ws[ws_], (wc0, w0 + 1, w1), [wc0 + 6, wc0 + 7], w0, "Picks per hour: current vs optimized",
                  "picks/h", colors=[SLATE, TEAL], w=15, h=7.5)
        ns, n0, n1, nc0 = self.ref["net_rank"]
        bar_chart(ws, "J23", self.ws[ns], (nc0 + 1, n0 + 1, n1), [nc0 + 9], n0, "MFC break-even (x today's parcel volume)",
                  "x volume", colors=[MARIGOLD], horizontal=True, w=16, h=7.5)
        # Daily demand sparklines
        sr = 38
        section(ws, sr, 2, "90-day demand trend", 6)
        dsheet = SHEETS[2]
        for k, (rng, _) in enumerate(self.spark.get(dsheet, [])):
            r = sr + 1 + k
            name = self.ws[dsheet].cell(int(self.spark[dsheet][k][1][1:]), 11).value
            ws.cell(r, 2, name).font = font(10, True)
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            self.spark.setdefault(s, []).append((rng, f"C{r}"))
            ws.row_dimensions[r].height = 22
        # Health components
        section(ws, sr, 9, "Health score components", 5)
        hc = pd.DataFrame([{"Component": k.replace("_", " ").capitalize(), "Score": v, "Weight": cc["health"]["weights"][k]}
                           for k, v in cc["health"]["components"].items()])
        r0h, r1h, c0h, _ = write_df(ws, hc, sr + 1, 9, {"Score": "dec1", "Weight": "pct"}, "HealthComponents", {})
        data_bar(ws, col_range(r0h + 1, r1h, c0h + 1), TEAL)
        ar = sr + 9
        al = self.p("cc_alerts.parquet").head(12)
        al["city"] = al["city_id"].map(self.names)
        al = al[["severity", "area", "city", "message"]]
        al.columns = ["Severity", "Area", "City", "Alert"]
        section(ws, ar, 2, "Top alerts (ranked by severity and impact)", 16)
        r0a, r1a, c0a, _ = write_df(ws, al, ar + 1, 2, None, "Alerts", {})
        ws.merge_cells(start_row=ar + 1, start_column=5, end_row=ar + 1, end_column=17)
        for r in range(r0a + 1, r1a + 1):
            ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=17)
        for sev, fill in (("high", "F8D9D6"), ("medium", "FFF6DD"), ("low", "EEF2F6"), ("info", "E3EEF0")):
            ws.conditional_formatting.add(col_range(r0a + 1, r1a, c0a), CellIsRule(operator="equal", formula=[f'"{sev}"'],
                                          fill=PatternFill("solid", fgColor=fill)))
        up = cc["upstream"]
        note(ws, r1a + 2, 2, "Portfolio links: " + ", ".join(f"{k} = {v}" for k, v in up.items())
             + f". Safety-stock linkage: {cc['safety_stock']['status']}.", 16)

    # ------------------------------------------------------------------ build
    def build(self) -> Path:
        for step in (self.master, self.depots, self.demand, self.routing, self.stops, self.fleet, self.slotting, self.wh_perf,
                     self.simulation, self.network, self.economics, self.ml, self.kpis, self.comparison, self.executive):
            step()
        for ws in self.wb.worksheets:
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.wb.calculation = CalcProperties(fullCalcOnLoad=True)
        self.wb.properties.title = "Logistics and Network Design"
        self.wb.properties.creator = "Kushagra Kartikey - Supply Chain Command Center"
        path = resolve_path(self.cfg, "workbook")
        self.wb.save(path)
        n = inject_sparklines(path, {t: i + 1 for i, t in enumerate(SHEETS)}, self.spark)
        log.info("Workbook written to %s (%d sheets, %d sparklines)", path, len(SHEETS), n)
        return path


def run(cfg: dict[str, Any]) -> Path:
    return WorkbookBuilder(cfg).build()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", default=None)
    run(load_config(parser.parse_args().config))


if __name__ == "__main__":
    main()
