"""Reporting layer: sparkline injection, Dash figure builders, NPV parity, payload structure."""
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.logistics.reporting.xl_helpers import inject_sparklines, kpi_card, write_df

ROOT = Path(__file__).resolve().parents[1]
PAYLOAD = ROOT / "dashboard" / "static" / "data" / "payload.json"


def test_sparkline_injection_produces_valid_workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Order & Demand"
    for i in range(1, 11):
        ws.cell(i, 1, i * i)
    kpi_card(ws, 1, 3, "Test", "=SUM(A1:A10)", "0")
    path = tmp_path / "t.xlsx"
    wb.save(path)
    n = inject_sparklines(path, {"Order & Demand": 1}, {"Order & Demand": [("'Order & Demand'!A1:A10", "E5")]})
    assert n == 1
    xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    assert "x14:sparkline" in xml and "&amp;" in xml
    assert load_workbook(path)["Order & Demand"]["C1"].value == "=SUM(A1:A10)"


def test_write_df_styles_and_table(tmp_path):
    import pandas as pd
    wb = Workbook()
    r0, r1, c0, c1 = write_df(wb.active, pd.DataFrame({"a": [1, 2], "b": [True, False]}), 3, 2, {"a": "int"}, "T1")
    assert (r0, r1, c0, c1) == (3, 5, 2, 3)
    assert wb.active["C4"].value == "Yes" and "T1" in wb.active.tables


@pytest.fixture(scope="module")
def payload():
    if not PAYLOAD.exists():
        pytest.skip("payload.json not built")
    return json.loads(PAYLOAD.read_text())


def test_payload_covers_all_cities_and_modules(payload):
    ids = [c["id"] for c in payload["meta"]["cities"]]
    assert ids == ["MUM", "DEL", "BLR", "CHE", "KOL"]
    for key in ("routing", "fleet", "warehouse", "simulation", "network", "demand", "alerts"):
        assert key in payload
    for c in ids:
        assert payload["routing"]["routes"][c]["optimized"]
        assert len(payload["warehouse"][c]["grid"]["optimized"]) == 20


def test_dash_figures_and_npv_parity(payload):
    import sys
    sys.path.insert(0, str(ROOT))
    from dashboard import app as dash_app
    city = "MUM"
    for fig in (dash_app.fig_cost_by_city(payload, city), dash_app.fig_route_map(payload, city, "optimized"),
                dash_app.fig_fleet(payload, city), dash_app.fig_heatmap(payload, city, "optimized"),
                dash_app.fig_sim(payload, city), dash_app.fig_network_map(payload, city, 2)):
        assert len(fig.data) > 0
    conf = next(c for c in payload["network"]["cities"][city]["configs"] if c["p"] == 1)
    v = dash_app.npv(conf, payload["meta"]["wacc"], payload["meta"]["growth"], payload["meta"]["horizon"])
    assert v == pytest.approx(conf["npv"], rel=1e-3, abs=50)   # client-side NPV == network_designer NPV
    assert dash_app.create_app(payload).layout is not None
